#!/usr/bin/env python3
"""Single-slide PPT: exact infographic only."""

from pathlib import Path
import shutil

from pptx import Presentation
from pptx.util import Inches

ASSETS = Path(__file__).resolve().parent.parent / "assets"
INFOGRAPHIC = ASSETS / "ai-industry-chain-infographic-zh.png"
SOURCE = Path("/opt/cursor/artifacts/assets/ai-industry-chain-infographic-zh.png")


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    if not INFOGRAPHIC.exists() and SOURCE.exists():
        shutil.copy2(SOURCE, INFOGRAPHIC)

    if not INFOGRAPHIC.exists():
        raise FileNotFoundError("Infographic PNG missing")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(INFOGRAPHIC), 0, 0, width=prs.slide_width, height=prs.slide_height)

    out = Path("/workspace/AI-industry-chain-infographic-v2.pptx")
    prs.save(str(out))
    shutil.copy2(out, Path("/workspace/output/AI产业链信息图-v2.pptx"))
    print(f"Saved: {out} ({out.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
