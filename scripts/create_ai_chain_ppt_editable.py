#!/usr/bin/env python3
"""Fully editable 16:9 AI industry chain PPT with embedded column images."""

from pathlib import Path
import shutil

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ASSETS = Path(__file__).resolve().parent.parent / "assets"
PNG_16x9 = ASSETS / "ai-industry-chain-infographic-zh-16x9.png"
PNG_SRC = ASSETS / "ai-industry-chain-infographic-zh.png"
CROPS = ASSETS / "column-crops-16x9"
SOURCE = Path("/opt/cursor/artifacts/assets/ai-industry-chain-infographic-zh.png")

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(10, 22, 40)
NAVY_CARD = RGBColor(14, 28, 50)
NAVY_LIGHT = RGBColor(18, 35, 60)
GOLD = RGBColor(245, 197, 24)
WHITE = RGBColor(255, 255, 255)
BLUE = RGBColor(59, 130, 246)
GRAY = RGBColor(180, 190, 210)
FONT = "Microsoft YaHei"

COLUMNS = [
    {"title": "算力层", "tagline": "生于硅片，GPU 与 HBM\n驱动一切。", "icon": "芯"},
    {"title": "基础设施层", "tagline": "电力、液冷与万卡\n集群规模化。", "icon": "云"},
    {"title": "模型层", "tagline": "数据训练，算力炼成\n认知能力。", "icon": "脑"},
    {"title": "智能体层", "tagline": "工具编排，自主规划\n与闭环执行。", "icon": "链"},
    {"title": "应用层", "tagline": "落地千行百业，\n嵌入日常工作。", "icon": "用"},
]
METRICS = ["资本开支", "推理延迟", "单位经济", "产品契合"]

# Column photo regions on 1920x1080 infographic (fractions)
COL_PHOTO_FRAC = [
    (0.025, 0.50, 0.195, 0.88),
    (0.210, 0.50, 0.385, 0.88),
    (0.400, 0.50, 0.575, 0.88),
    (0.590, 0.50, 0.765, 0.88),
    (0.780, 0.50, 0.975, 0.88),
]


def set_fill(shape, color):
    f = shape.fill
    f.solid()
    f.fore_color.rgb = color


def textbox(slide, l, t, w, h, text, size=14, bold=False, color=WHITE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def ensure_png_16x9() -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    if not PNG_SRC.exists() and SOURCE.exists():
        shutil.copy2(SOURCE, PNG_SRC)
    if PNG_16x9.exists():
        return PNG_16x9
    img = Image.open(PNG_SRC)
    tw, th = 1920, 1080
    sw, sh = img.size
    scale = max(tw / sw, th / sh)
    nw, nh = int(sw * scale), int(sh * scale)
    resized = img.resize((nw, nh), Image.Resampling.LANCZOS)
    left = (nw - tw) // 2
    top = (nh - th) // 2
    out = resized.crop((left, top, left + tw, top + th))
    out.save(PNG_16x9)
    return PNG_16x9


def crop_column_images() -> list[Path]:
    CROPS.mkdir(parents=True, exist_ok=True)
    img = Image.open(ensure_png_16x9())
    w, h = img.size
    paths = []
    for i, (l, t, r, b) in enumerate(COL_PHOTO_FRAC):
        p = CROPS / f"col-{i + 1}.png"
        img.crop((int(l * w), int(t * h), int(r * w), int(b * h))).save(p)
        paths.append(p)
    return paths


def build_editable_slide(prs, column_images: list[Path]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sw, sh = prs.slide_width, prs.slide_height

    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, sw, sh)
    set_fill(bg, NAVY)
    bg.line.fill.background()

    hdr_h = Inches(1.55)
    hdr = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, sw, hdr_h)
    set_fill(hdr, NAVY_LIGHT)
    hdr.line.fill.background()

    textbox(slide, Inches(0.45), Inches(0.28), Inches(12.2), Inches(0.65),
            "做 AI 投资？最大的差异不是模型，而是产业链。", 28, True)
    textbox(slide, Inches(0.45), Inches(0.95), Inches(12.2), Inches(0.4),
            "五层结构，一条链条，价值从芯片流向应用。", 16, color=GRAY)

    col_top = Inches(1.75)
    col_h = Inches(4.35)
    mx = Inches(0.35)
    gap = Inches(0.18)
    col_w = (sw - 2 * mx - 4 * gap) // 5

    for i, col in enumerate(COLUMNS):
        left = mx + i * (col_w + gap)

        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, col_top, col_w, col_h)
        set_fill(card, NAVY_CARD)
        card.line.color.rgb = RGBColor(30, 55, 90)

        textbox(slide, left + Inches(0.1), col_top + Inches(0.12), col_w - Inches(0.2), Inches(0.42),
                col["title"], 20, True, align=PP_ALIGN.CENTER)
        textbox(slide, left + Inches(0.08), col_top + Inches(0.58), col_w - Inches(0.16), Inches(0.72),
                col["tagline"], 12, color=GRAY, align=PP_ALIGN.CENTER)

        isz = Inches(0.68)
        il = left + (col_w - isz) // 2
        it = col_top + Inches(1.48)
        ico = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, il, it, isz, isz)
        set_fill(ico, BLUE)
        ico.line.fill.background()
        textbox(slide, il, it + Inches(0.1), isz, isz - Inches(0.2), col["icon"], 22, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        # Embedded photo from infographic (editable: delete/replace in PPT)
        pt = col_top + Inches(2.38)
        ph = Inches(1.72)
        pl = left + Inches(0.12)
        pw = col_w - Inches(0.24)
        if column_images[i].exists():
            slide.shapes.add_picture(str(column_images[i]), pl, pt, width=pw, height=ph)

    ft = Inches(6.35)
    fh = Inches(1.05)
    foot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, ft, sw, fh)
    set_fill(foot, NAVY_LIGHT)
    foot.line.fill.background()

    mb = slide.shapes.add_textbox(Inches(0.45), ft + Inches(0.28), Inches(7.0), Inches(0.55))
    tf = mb.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = "同一技术，不同层级。"
    r1.font.name, r1.font.size, r1.font.color.rgb = FONT, Pt(15), WHITE
    r2 = p.add_run()
    r2.text = "看懂产业链，才能抓住价值。"
    r2.font.name, r2.font.size, r2.font.bold, r2.font.color.rgb = FONT, Pt(15), True, GOLD

    mw = Inches(1.35)
    ms = Inches(8.1)
    mg = Inches(0.25)
    for j, label in enumerate(METRICS):
        x = ms + j * (mw + mg)
        y = ft + Inches(0.18)
        c = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x, y, Inches(0.42), Inches(0.42))
        set_fill(c, RGBColor(30, 55, 90))
        c.line.color.rgb = BLUE
        textbox(slide, x - Inches(0.05), y + Inches(0.46), mw, Inches(0.35), label, 11, align=PP_ALIGN.CENTER)


def build_reference_slide(prs, png: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    set_fill(bg, NAVY)
    bg.line.fill.background()
    textbox(slide, Inches(0.4), Inches(0.15), Inches(8), Inches(0.35),
            "参考：完整信息图（第 1 页为可编辑版）", 13, color=GRAY)
    slide.shapes.add_picture(str(png), Inches(0.15), Inches(0.5), width=Inches(13.0))


def main():
    png = ensure_png_16x9()
    crops = crop_column_images()

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    build_editable_slide(prs, crops)
    build_reference_slide(prs, png)

    out = Path("/workspace/AI-industry-chain-editable-16x9.pptx")
    prs.save(str(out))
    shutil.copy2(out, Path("/workspace/output/AI产业链信息图-可编辑-16x9.pptx"))
    print(f"Saved: {out}")
    print("  Slide 1: editable text + shapes + 5 column images (16:9)")
    print("  Slide 2: full infographic reference")


if __name__ == "__main__":
    main()
